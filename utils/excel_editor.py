import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter, column_index_from_string


class ExcelEditor:
    """Utility class for reading and editing Excel (.xlsx) files."""

    def __init__(self, filepath: str):
        self.filepath = filepath
        try:
            self.workbook = openpyxl.load_workbook(filepath)
        except FileNotFoundError:
            self.workbook = Workbook()
        self.sheet = self.workbook.active

    # ------------------------------------------------------------------
    # Sheet helpers
    # ------------------------------------------------------------------

    def select_sheet(self, name: str) -> None:
        """Switch the active sheet by name."""
        if name not in self.workbook.sheetnames:
            raise ValueError(f"Sheet '{name}' not found. Available: {self.workbook.sheetnames}")
        self.sheet = self.workbook[name]

    def create_sheet(self, name: str) -> None:
        """Create a new sheet and make it active."""
        self.sheet = self.workbook.create_sheet(title=name)

    @property
    def sheet_names(self) -> list:
        return self.workbook.sheetnames

    # ------------------------------------------------------------------
    # Cell operations
    # ------------------------------------------------------------------

    def read_cell(self, row: int, col: int | str) -> object:
        """Read a cell value. col can be an int or a letter string (e.g. 'A')."""
        col = self._normalize_col(col)
        return self.sheet.cell(row=row, column=col).value

    def write_cell(self, row: int, col: int | str, value: object) -> None:
        """Write a value to a cell."""
        col = self._normalize_col(col)
        self.sheet.cell(row=row, column=col, value=value)

    # ------------------------------------------------------------------
    # Row operations
    # ------------------------------------------------------------------

    def read_row(self, row: int) -> list:
        """Return all values in a row as a list."""
        return [cell.value for cell in self.sheet[row]]

    def write_row(self, row: int, values: list) -> None:
        """Write a list of values into a row, starting at column 1."""
        for col_idx, value in enumerate(values, start=1):
            self.sheet.cell(row=row, column=col_idx, value=value)

    def append_row(self, values: list) -> None:
        """Append a row of values after the last used row."""
        self.sheet.append(values)

    def delete_row(self, row: int) -> None:
        """Delete a row (1-indexed) and shift remaining rows up."""
        self.sheet.delete_rows(row)

    def insert_row(self, row: int, values: list | None = None) -> None:
        """Insert an empty row at the given position, then optionally fill it."""
        self.sheet.insert_rows(row)
        if values:
            self.write_row(row, values)

    # ------------------------------------------------------------------
    # Column operations
    # ------------------------------------------------------------------

    def read_column(self, col: int | str) -> list:
        """Return all values in a column as a list."""
        col = self._normalize_col(col)
        return [cell.value for cell in self.sheet[get_column_letter(col)]]

    def delete_column(self, col: int | str) -> None:
        """Delete a column and shift remaining columns left."""
        col = self._normalize_col(col)
        self.sheet.delete_cols(col)

    # ------------------------------------------------------------------
    # Full-sheet operations
    # ------------------------------------------------------------------

    def read_all(self) -> list[list]:
        """Return all sheet data as a list of rows (each row is a list)."""
        return [[cell.value for cell in row] for row in self.sheet.iter_rows()]

    def get_headers(self) -> list:
        """Return the first row (assumed to be headers)."""
        return self.read_row(1)

    def find_row_by_value(self, col: int | str, value: object) -> int | None:
        """Find the first row number where the given column equals value."""
        col = self._normalize_col(col)
        for row in self.sheet.iter_rows(min_col=col, max_col=col):
            cell = row[0]
            if cell.value == value:
                return cell.row
        return None

    def update_row_by_value(self, search_col: int | str, search_value: object, updates: dict) -> bool:
        """
        Find a row where search_col == search_value and apply column->value updates.

        updates is a dict mapping column identifiers (int or letter) to new values.
        Returns True if a row was found and updated, False otherwise.
        """
        row_num = self.find_row_by_value(search_col, search_value)
        if row_num is None:
            return False
        for col, value in updates.items():
            self.write_cell(row_num, col, value)
        return True

    # ------------------------------------------------------------------
    # Persistence
    # ------------------------------------------------------------------

    def save(self, filepath: str | None = None) -> None:
        """Save the workbook. Uses the original path when no path is given."""
        self.workbook.save(filepath or self.filepath)

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _normalize_col(col) -> int:
        """Convert a column letter ('A') or int to a 1-based column index."""
        if isinstance(col, str):
            return column_index_from_string(col)
        return col
